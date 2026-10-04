import hashlib
import json
import shutil
import subprocess
from pathlib import Path
from datetime import datetime, timezone
from docx import Document

root = Path(__file__).resolve().parents[1]
dest = root / 'evidencias/revision-hallazgos-20260930'
dest.mkdir(exist_ok=False)
files = list(root.glob('__tests__/*.test.ts')) + list(root.glob('utils/*.ts'))
files += [root / p for p in ['ai-engine/llm_service.py', 'ai-engine/main.py', 'ai-engine/test_unit_llm.py', 'services/paymentSimulation.ts', 'components/TriageEvaluation.tsx', 'app/(patient)/payment.tsx', 'app/(patient)/waiting-room.tsx', 'app/(doctor)/dashboard.tsx', 'app/(doctor)/medical-record.tsx', 'app/(admin)/reports.tsx', 'jest.config.js', 'package.json', 'scripts/record-unit-tests.cjs']]
for src in files:
    out = dest / 'antes' / src.relative_to(root)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, out)
preserve = list((root / 'evidencias/pruebas-completas').rglob('*')) + list((root / 'evidencias/pruebas-usuarios').rglob('*')) + list((root / 'output/documents').glob('*.docx'))
manifest = [{'archivo':str(p.relative_to(root)), 'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in preserve if p.is_file()]
(dest / 'historico-preservado.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
metadata = {'inicioUTC': datetime.now(timezone.utc).isoformat(), 'commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=root,text=True).strip(), 'estadoInicial':subprocess.check_output(['git','status','--short'],cwd=root,text=True), 'documentosAcademicos': 'Sin modificar'}
(dest / 'inicio.json').write_text(json.dumps(metadata, ensure_ascii=False,indent=2),encoding='utf-8')
for label, path in [('requisitos',Path(r'C:\Users\inutil\Downloads\Documento_Oficial_V2 (5).docx')),('informe-anterior',root/'output/documents/Informe_detallado_pruebas_unitarias_telemedicina.docx')]:
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for i,table in enumerate(doc.tables):
        lines.append(f'TABLA {i+1}')
        lines.extend(' | '.join(c.text for c in row.cells) for row in table.rows)
    (dest / f'{label}.txt').write_text('\n'.join(lines),encoding='utf-8')
print(dest)
