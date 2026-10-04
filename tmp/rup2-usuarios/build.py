from pathlib import Path
from copy import deepcopy
import json, re, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
E=ROOT/'evidencias/pruebas-usuarios'
REF=Path(r'C:\Users\inutil\Downloads\Documento_Oficial_V2 (5).docx')
r=json.loads((E/'final.json').read_text(encoding='utf-8'))
assert [r[k] for k in ['numTotalTests','numPassedTests','numFailedTests','numPendingTests','numTotalTestSuites']]==[30,30,0,0,2]
assert 'Tests:       30 passed, 30 total' in (E/'final.log').read_text(encoding='utf-8-sig')
files=['utils/validation.ts','utils/userValidation.ts','__tests__/loginValidation.test.ts','__tests__/userValidation.test.ts','jest.config.js']
hashes={f:hashlib.sha256((ROOT/f).read_bytes()).hexdigest() for f in files}
for f in files: assert (ROOT/f).read_bytes()==(E/'fuentes'/(f+'.txt')).read_bytes()
actual={}
for suite in r['testResults']:
 for i,a in enumerate(suite['assertionResults']):
  m=re.search(r'PU-\d+',a['title']); key=m.group() if m else f'PU-{i+1:02d}'
  assert a['status']=='passed'; actual[key]=a['title']
assert len(actual)==30
ref=Document(REF)
levels=[0]*4
section=None
for para in ref.paragraphs:
 if not para.style.name.startswith('Heading '):continue
 n=int(para.style.name.split()[-1])-1
 if n>3:continue
 x=para._p.find('w:pPr/w:numPr/w:numId',para._p.nsmap)
 if x is not None and x.get(qn('w:val'))=='0':continue
 levels[n]+=1
 for j in range(n+1,4):levels[j]=0
 if levels[0]==3 and para.text.strip()=='PRUEBAS UNITARIAS':section='.'.join(map(str,levels[:n+1]))
assert section

d=Document()
s=d.sections[0]
s.page_width=Cm(21.59);s.page_height=Cm(27.94)
s.left_margin=Cm(3);s.right_margin=Cm(2);s.top_margin=Cm(2.5);s.bottom_margin=Cm(2.5)
for style in d.styles:
 if style.type==1:
  style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 st=d.styles[name];st.font.size=Pt(12);st.paragraph_format.space_after=Pt(8);st.paragraph_format.line_spacing=1.15
 st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 st.paragraph_format.keep_together=True
for name in ['Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.bold=True;st.paragraph_format.keep_with_next=True
 st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
d.core_properties.title='Pruebas unitarias de validación de usuarios'
d.core_properties.author='';d.core_properties.comments=''

def p(t,style=None):return d.add_paragraph(t,style)
def source():
 q=p('Fuente: Elaboración propia, 2026.');q.paragraph_format.space_after=Pt(10)
 q.runs[0].font.size=Pt(10)
def caption(t):
 q=p(t);q.paragraph_format.keep_with_next=True
 q.runs[0].bold=True
def table(rows,widths,header=True,summary=False):
 t=d.add_table(rows=0,cols=len(widths));t.autofit=False;t.style='Table Grid'
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for i,values in enumerate(rows):
  row=t.add_row()
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i==0 and header:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for j,(cell,text,w) in enumerate(zip(row.cells,values,widths)):
   cell.width=Cm(w);cell.text=text;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
   tc.append(margins)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6B8AF' if not summary else 'EAC4C4');tc.append(shade)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
    para.paragraph_format.line_spacing=1.1
    para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or j==0 or (len(widths)==5 and j==4) else WD_ALIGN_PARAGRAPH.LEFT
    for run in para.runs:run.font.size=Pt(11);run.bold=(i==0)
 for cell in t.rows[-1].cells:
  for para in cell.paragraphs:para.paragraph_format.keep_with_next=True
 return t

p(section+'. PRUEBAS UNITARIAS DE VALIDACIÓN DE USUARIOS','Title')
p('Se realizaron pruebas unitarias para comprobar las validaciones locales del inicio de sesión, el registro y la edición de perfiles de pacientes y médicos. Se evaluaron cuatro funciones de forma aislada mediante 30 casos automatizados con datos ficticios. Estas pruebas verifican las respuestas de la validación; no comprueban el acceso a Supabase ni el funcionamiento completo del módulo de usuarios.')
caption('Tabla. Herramientas utilizadas en las pruebas unitarias.')
table([['Herramienta','Versión','Función en las pruebas'],['Jest','29.7.0','Ejecuta los casos y compara las respuestas con lo esperado.'],['jest-expo','57.0.5','Adapta la configuración de las pruebas al proyecto Expo.'],['Node.js','24.15.0','Permite ejecutar las herramientas de prueba.'],['npm','11.12.1','Inicia la ejecución mediante el script de pruebas del proyecto.']],[3,2.3,11.29])
source()
caption('Tabla. Resumen de pruebas unitarias de validación de usuarios.')
table([
 ['Objetivo','Comprobar las reglas locales de validación de datos de usuarios en los escenarios definidos.'],
 ['Acción','Se ejecutaron 30 casos con Jest: nueve de inicio de sesión, doce de registro, cuatro de perfil del paciente y cinco de perfil del médico. En cada caso se prepararon datos ficticios, se llamó a la función y se comparó su respuesta con el resultado esperado.'],
 ['Efecto','Los 30 casos aprobaron sus comprobaciones en dos conjuntos de pruebas, sin casos fallidos ni omitidos. Dos casos reprodujeron limitaciones relacionadas con datos formados solo por espacios, que siguen pendientes de corrección. Véase el apéndice «Pruebas unitarias de validación de usuarios».']
 ],[2.55,14.04],header=False,summary=True)
source()

d.add_page_break()
q=p('APÉNDICE','Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
q=p('PRUEBAS UNITARIAS DE VALIDACIÓN DE USUARIOS','Title');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
p('Las siguientes tablas presentan los 30 casos ejecutados con Jest, agrupados según la validación evaluada. Los pasos describen llamadas automatizadas a las funciones, no recorridos manuales en las pantallas. La palabra «Éxito» indica que la respuesta coincidió con la expectativa del caso; no significa que el sistema esté libre de problemas.')

data=[]
def add(scenario,prepare,expected):
 data.append((scenario,'1. '+prepare+'\n2. Ejecutar la función de validación.\n3. Comparar la respuesta con la esperada.',expected))
add('Ambos campos del inicio de sesión están vacíos.','Preparar el correo y la contraseña como cadenas vacías.','La validación rechaza los datos y solicita ingresar el correo y la contraseña.')
add('La contraseña está vacía.','Preparar un correo ficticio completo y una contraseña vacía.','La validación rechaza los datos y solicita ingresar la contraseña.')
add('El correo está vacío.','Preparar un correo vacío y una contraseña ficticia completa.','La validación rechaza los datos y solicita ingresar el correo electrónico.')
add('El correo contiene únicamente espacios.','Preparar un correo de cinco espacios y una contraseña ficticia completa.','La validación rechaza el correo de solo espacios y solicita ingresar el correo electrónico.')
add('El correo y la contraseña están completos.','Preparar un correo y una contraseña ficticios, ambos con contenido.','Los datos superan la validación local y no se devuelve un mensaje de error. No se comprueba la autenticación.')
add('Los dos argumentos se reciben sin un valor definido.','Preparar ambos argumentos con el valor undefined.','La validación rechaza las entradas y solicita ingresar el correo y la contraseña.')
add('Los dos argumentos se reciben con valor nulo.','Preparar ambos argumentos con el valor null.','La validación rechaza las entradas y solicita ingresar el correo y la contraseña.')
add('El correo tiene espacios al principio y al final.','Preparar un correo ficticio rodeado de espacios y una contraseña con contenido.','Los datos superan la validación local y no se devuelve un mensaje de error.')
add('La contraseña contiene únicamente espacios.','Preparar un correo ficticio completo y una contraseña de tres espacios.','Los datos superan la validación local porque la contraseña tiene caracteres. No se recortan sus espacios ni se comprueba la autenticación.')
for field in ['correo','contraseña','nombre completo','carnet de identidad','teléfono','dirección']:
 add('El '+field+' está vacío.' if field not in ['contraseña','dirección'] else 'La '+field+' está vacía.',
     'Preparar un registro de paciente con datos ficticios completos y dejar vacío el campo de '+field+'.',
     'La validación rechaza el registro y solicita completar todos los campos comunes.')
add('El paciente completa sus datos sin título ni especialidad.','Preparar todos los campos comunes del paciente, sin especialidad y sin documento de título.','El registro del paciente supera la validación local sin requerir datos exclusivos del médico.')
add('El médico no tiene una especialidad seleccionada.','Preparar los campos comunes del médico y un documento ficticio; dejar vacía la especialidad.','La validación rechaza los datos y solicita seleccionar una especialidad y adjuntar el título.')
add('El médico no presenta un documento de título.','Preparar los campos comunes y la especialidad del médico; establecer el documento como nulo.','La validación rechaza los datos y solicita seleccionar una especialidad y adjuntar el título.')
add('El médico tiene completos sus datos y requisitos.','Preparar los campos comunes, una especialidad y un objeto de documento ficticio.','El registro médico supera la validación local. No se crea una cuenta ni se carga el documento.')
add('Faltan un campo común y el título del médico.','Preparar datos de médico con el correo vacío y el documento nulo.','La validación rechaza los datos y muestra primero la solicitud de completar los campos comunes.')
add('El nombre del paciente contiene únicamente espacios.','Preparar el registro de paciente con datos ficticios completos y sustituir el nombre por tres espacios.','El nombre de solo espacios supera la validación local y no genera un mensaje de error. Se reproduce el comportamiento existente.')
add('El teléfono del paciente está vacío.','Preparar un teléfono vacío y la dirección «Dirección ficticia».','La validación rechaza los datos y solicita completar el teléfono y la dirección.')
add('La dirección del paciente está vacía.','Preparar el teléfono ficticio «00000000» y una dirección vacía.','La validación rechaza los datos y solicita completar el teléfono y la dirección.')
add('El teléfono y la dirección del paciente tienen contenido.','Preparar el teléfono «00000000» y la dirección «Dirección ficticia».','Los datos superan la validación local y no se devuelve un mensaje de error.')
add('El teléfono y la dirección contienen únicamente espacios.','Preparar el teléfono y la dirección con tres espacios en cada campo.','Ambos campos superan la validación local y no generan un mensaje de error. Se reproduce el comportamiento existente.')
add('El nombre del médico está vacío.','Preparar un nombre vacío y el celular ficticio «00000000».','La validación rechaza los datos e informa que el nombre y el celular no pueden estar vacíos.')
add('El celular del médico está vacío.','Preparar el nombre «Médico de Prueba» y un celular vacío.','La validación rechaza los datos e informa que el nombre y el celular no pueden estar vacíos.')
add('El nombre del médico contiene únicamente espacios.','Preparar un nombre de tres espacios y el celular ficticio «00000000».','La validación rechaza el nombre de solo espacios e informa que el nombre y el celular no pueden estar vacíos.')
add('El celular del médico contiene únicamente espacios.','Preparar el nombre «Médico de Prueba» y un celular de tres espacios.','La validación rechaza el celular de solo espacios e informa que el nombre y el celular no pueden estar vacíos.')
add('El nombre y el celular tienen espacios exteriores.','Preparar « Médico de Prueba » y « 00000000 », conservando los espacios exteriores.','Los datos superan la validación local porque ambos campos contienen información además de los espacios.')
assert len(data)==30
groups=[('Inicio de sesión','validateLoginFields',1,9),('Registro de usuarios','validateRegistrationFields',10,21),('Edición del perfil del paciente','validatePatientProfileFields',22,25),('Edición del perfil del médico','validateDoctorProfileFields',26,30)]
for title,fn,start,end in groups:
 # Los grupos continúan sin saltos forzados.
 caption('Validación evaluada: '+title+'.')
 q=p('Se ejecutaron '+str(end-start+1)+' casos sobre la función '+fn+'.');q.paragraph_format.keep_with_next=True
 if start==10:p('Los datos comunes del registro son correo, contraseña, nombre completo, carnet de identidad, teléfono y dirección. En cada prueba se modifica únicamente lo indicado; los demás campos conservan datos ficticios completos. El documento del médico es un objeto de prueba y no se abre ni se envía a un servicio.')
 rows=[['Caso de prueba','Escenario de prueba','Pasos de la prueba','Resultado esperado','Éxito / Fracaso']]
 for n in range(start,end+1):
  scenario,steps,expected=data[n-1]
  status='Éxito' if n not in [21,25] else 'Éxito al reproducir el comportamiento actual; se detectó una limitación funcional'
  rows.append([f'PU-{n:02d}',scenario,steps,expected,status])
 table(rows,[1.7,2.9,5,3.89,3.1])
 source()
 if end==21:
  p('Hallazgo H-01. El caso PU-21 confirmó que un nombre formado solo por espacios supera la validación del registro. La prueba tuvo éxito al reproducir esa respuesta, pero el comportamiento constituye una limitación funcional que sigue pendiente de corrección.')
 if end==25:
  p('Hallazgo H-02. El caso PU-25 confirmó que un teléfono y una dirección formados solo por espacios superan la validación del perfil del paciente. La prueba tuvo éxito al reproducir esa respuesta, pero el comportamiento constituye una limitación funcional que sigue pendiente de corrección.')

out=ROOT/'output/documents/Pruebas_unitarias_usuarios_formato_RUP2.docx'
d.save(out)
audit={'heading':section,'counts':[9,12,4,5],'sources':hashes,'reference_hash':hashlib.sha256(REF.read_bytes()).hexdigest(),'cases':actual,'screenshots':[],'status':'Borrador pendiente de capturas reales; no se insertaron marcadores gráficos en el texto académico.'}
(ROOT/'tmp/rup2-usuarios/audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(out)
