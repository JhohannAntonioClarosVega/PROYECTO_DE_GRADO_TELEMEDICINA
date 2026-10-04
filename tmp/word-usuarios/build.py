from pathlib import Path
import json, re, hashlib
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
E = ROOT / 'evidencias/pruebas-usuarios'
OUT = ROOT / 'output/documents/Pruebas_unitarias_usuarios_capitulo_y_apendice.docx'
res = json.loads((E/'final.json').read_text(encoding='utf-8'))
env = json.loads((E/'entorno.json').read_text(encoding='utf-8'))
assert (res['numTotalTests'],res['numPassedTests'],res['numFailedTests'],res['numPendingTests'],res['numTotalTestSuites']) == (30,30,0,0,2)
assert 'Tests:       30 passed, 30 total' in (E/'final.log').read_text(encoding='utf-8-sig')
sources=['utils/validation.ts','utils/userValidation.ts','__tests__/loginValidation.test.ts','__tests__/userValidation.test.ts','jest.config.js']
for f in sources:
    assert (ROOT/f).read_bytes() == (E/'fuentes'/(f+'.txt')).read_bytes()
cases={}
for s in res['testResults']:
    for i,a in enumerate(s['assertionResults']):
        match=re.search(r'PU-\d+',a['title'])
        key=match.group() if match else f'PU-{i+1:02d}'
        assert a['status']=='passed'
        cases[key]=a['title']
assert len(cases)==30

d=Document()
s=d.sections[0]
s.page_width=Cm(21.59); s.page_height=Cm(27.94)
s.left_margin=Cm(3); s.right_margin=Cm(2); s.top_margin=Cm(2.5); s.bottom_margin=Cm(2.5)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
    style=d.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(12 if name not in ['Caption'] else 10)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.15
d.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Title','Heading 1','Heading 2','Heading 3']:
    d.styles[name].font.bold=True
    d.styles[name].paragraph_format.keep_with_next=True
    d.styles[name].paragraph_format.space_before=Pt(10)
d.styles['Title'].font.size=Pt(14)
for border in list(d.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
d.core_properties.title='Pruebas unitarias de validación de usuarios'
d.core_properties.subject='Contenido para el Capítulo III y apéndice de pruebas unitarias'
d.core_properties.author=''
d.core_properties.comments=''

def p(text='',style=None):
    return d.add_paragraph(text,style)
def head(text,level=2):
    return d.add_heading(text,level)
def source(text):
    q=p('Fuente: '+text,'Caption'); q.paragraph_format.keep_with_next=False
    for r in q.runs:r.bold=False
def caption(text):
    q=p(text,'Caption'); q.paragraph_format.keep_with_next=True
    for r in q.runs:r.bold=True
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    t.style='Table Grid'
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for c,x,w in zip(t.rows[0].cells,headers,widths): c.text=x; c.width=Cm(w)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,x,w in zip(cells,row,widths): c.text=str(x); c.width=Cm(w)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            tc=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'65'); el.set(qn('w:type'),'dxa'); mar.append(el)
            tc.append(mar)
            if i==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAD0CE'); tc.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.line_spacing=1.05
                para.paragraph_format.space_after=Pt(3)
                para.paragraph_format.space_before=Pt(2)
                for r in para.runs:r.font.size=Pt(10);r.bold=(i==0)
    return t
def figure_slot(title,explanation):
    caption(title)
    q=p('Evidencia gráfica pendiente de incorporación.');q.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.keep_with_next=True
    q.runs[0].italic=True
    q=p(explanation);q.paragraph_format.keep_with_next=True
    source('captura del proyecto; pendiente de incorporación.')

p('3.3.1. PRUEBAS UNITARIAS DE VALIDACIÓN DE USUARIOS','Title')
p('Se evaluaron cuatro funciones que verifican entradas antes de continuar con operaciones de la aplicación: validateLoginFields, validateRegistrationFields, validatePatientProfileFields y validateDoctorProfileFields. Las pruebas permiten comprobar las reglas locales de campos obligatorios, los requisitos adicionales del registro médico y el tratamiento de entradas vacías o compuestas por espacios.')
p('La ejecución comprendió 30 casos en dos suites de Jest: nueve de inicio de sesión, doce de registro, cuatro de perfil de paciente y cinco de perfil médico. Se invocaron las funciones con datos ficticios, sin autenticar cuentas ni consultar Supabase. Las funciones de registro y perfiles habían sido extraídas de las pantallas conservando sus condiciones; las pruebas utilizan la misma lógica que consume la aplicación.')
caption('Tabla. Herramientas de prueba y tecnologías del entorno')
table(['Categoría','Herramienta o tecnología','Versión','Función'],[
('Prueba','Jest','29.7.0','Ejecuta los casos y compara resultados mediante aserciones.'),
('Prueba','jest-expo','57.0.5','Preset de configuración existente para las pruebas en Expo.'),
('Apoyo de prueba','@types/jest','30.0.0','Definiciones de tipos de las funciones de Jest.'),
('Ejecución','Node.js / npm','24.15.0 / 11.12.1','Ejecuta Jest y el script npm test.'),
('Comprobación estática','TypeScript','6.0.3','Verificación adicional con tsc --noEmit; no sustituye a Jest.'),
('Aplicación','Expo','57.0.19','Entorno de la aplicación; no constituye una prueba de dispositivos.'),
('Aplicación','React / React Native','19.2.3 / 0.86.3','Tecnologías de interfaz; no se evaluaron sus pantallas en esta ejecución.')
],[2.7,3.1,2.7,8.09])
source('elaboración propia a partir de entorno.json y la configuración conservada, 2026.')
d.add_page_break()
caption('Tabla. Resumen de pruebas unitarias de validación de usuarios')
table(['Campo','Descripción'],[
('Objetivo','Verificar las reglas locales de las cuatro funciones de validación de usuarios en los escenarios definidos.'),
('Acción','Ejecutar 30 casos con Jest en loginValidation.test.ts y userValidation.test.ts. Preparar entradas ficticias, invocar las funciones y comparar isValid y error con los valores esperados.'),
('Efecto','Los 30 casos coincidieron con sus expectativas: dos suites aprobadas, cero casos fallidos y cero omitidos. PU-21 y PU-25 aprobaron como caracterización de limitaciones funcionales pendientes. El detalle consta en el apéndice «Pruebas unitarias de validación de usuarios».')
],[2.6,13.99])
source('elaboración propia a partir de final.json y final.log, ejecución del 30 de septiembre de 2026.')
figure_slot('Figura. Configuración de Jest y versiones del entorno','La evidencia corresponde a las versiones instaladas y al preset jest-expo declarado en jest.config.js.')
figure_slot('Figura. Ejecución de las pruebas unitarias de validación de usuarios','La salida de Jest identifica el comando, las dos suites y sus resultados. La ejecución documentada en final.json y final.log comenzó el 30 de septiembre de 2026 a las 14:10:00, hora de Bolivia, y reportó 30 casos aprobados.')
p('La aprobación en Jest significa que el resultado coincidió con la expectativa de cada caso. No equivale a afirmar que todo comportamiento sea funcionalmente adecuado. H-01 identifica la aceptación de un nombre de solo espacios en el registro; H-02 identifica la aceptación de teléfono y dirección de solo espacios en el perfil de paciente. Ambas limitaciones permanecen pendientes. La evaluación no comprende autenticación real, persistencia, autorización por roles, recorridos de interfaz ni otros módulos.')

d.add_page_break()
p('APÉNDICE','Title')
p('Pruebas unitarias de validación de usuarios','Title')
p('Este apéndice presenta las entradas, los pasos automatizados, las expectativas y los resultados de los 30 casos ejecutados. Los casos se agrupan por función. Los estados se sustentan en las aserciones registradas por Jest y no representan comprobaciones manuales en las pantallas de la aplicación.')
head('Datos de ejecución')
when=datetime.fromtimestamp(res['startTime']/1000,timezone(timedelta(hours=-4)))
p('Fecha y hora de inicio: '+when.strftime('%d/%m/%Y, %H:%M:%S')+' (America/La_Paz, UTC−04:00). Entorno: Windows y PowerShell, con Node.js 24.15.0 y npm 11.12.1. Resultado: 30 aprobados, 0 fallidos y 0 omitidos, distribuidos en dos suites. Jest informó 0 snapshot tests.')
p('Comando exacto registrado:')
q=p(env['finalCommand']);q.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
for r in q.runs:r.font.name='Consolas';r.font.size=Pt(10)
p('Archivos ejecutados: __tests__/loginValidation.test.ts y __tests__/userValidation.test.ts. Funciones evaluadas: utils/validation.ts y utils/userValidation.ts. La suite original contiene PU-01 a PU-05; la segunda contiene PU-06 a PU-30. Los identificadores de los primeros cinco casos se asignan según su orden en el archivo original.')
p('Referencia del proyecto: commit '+env['commit']+'. Existían cambios locales al ejecutar las pruebas, incluidos los archivos de validación y pruebas; por tanto, el commit aislado no reproduce el estado evaluado. El inventario consta en entorno.json y las copias de los archivos evaluados se conservan en fuentes/.')
p('Se utilizaron datos ficticios y un objeto de documento ficticio. No se abrió ni subió ese documento y no se realizaron llamadas a servicios externos. La comprobación adicional de TypeScript terminó con código de salida 0 y sin diagnósticos; se identifica en entorno.json y typescript.log.')
head('Identificación de evidencias')
p('Las rutas siguientes son relativas a evidencias/pruebas-usuarios/: E1 = final.json, resultados estructurados de Jest; E2 = final.log, salida textual; E3 = fuentes/__tests__/loginValidation.test.ts.txt; E4 = fuentes/__tests__/userValidation.test.ts.txt; E5 = entorno.json; E6 = fuentes/utils/validation.ts.txt y fuentes/utils/userValidation.ts.txt. usos-en-pantallas.txt identifica las llamadas desde la aplicación. Los archivos E1 y E2 respaldan los estados; E3, E4 y E6 permiten revisar entradas y aserciones.')
head('Convenciones de las tablas')
p('V representa isValid=true y ausencia de error. F representa isValid=false acompañado del mensaje indicado. M1: «Por favor ingresa tu correo y contraseña.»; M2: «Por favor ingresa tu contraseña.»; M3: «Por favor ingresa tu correo electrónico.»; M4: «Por favor completa todos los campos comunes.»; M5: «Por favor selecciona una especialidad y adjunta tu título.»; M6: «Por favor completa el número de teléfono y la dirección.»; M7: «El nombre y el celular no pueden estar vacíos».')
p('Pasos A: 1) preparar las entradas; 2) invocar la función; 3) comprobar isValid con toBe y el mensaje con toBe/toBeDefined, o su ausencia con toBeUndefined. Pasos B: 1) preparar las entradas; 2) invocar la función; 3) comparar el objeto completo con toEqual. Estos pasos corresponden a las instrucciones efectivamente ejecutadas por Jest.')
p('La columna «Resultado obtenido» expresa lo confirmado por las aserciones aprobadas; no es una transcripción de un objeto impreso en consola. «Aprobado en Jest» indica coincidencia con la expectativa. En PU-21 y PU-25 se distingue expresamente la caracterización de una limitación pendiente.')

data=[
('Ambos campos vacíos','email=""; password=""','F; M1'),
('Contraseña vacía','email=C1; password=""','F; M2'),
('Correo vacío','email=""; password=K1','F; M3'),
('Correo de solo espacios','email=5 espacios; password=K1','F; M3'),
('Campos completos','email=C2; password=K2','V'),
('Argumentos undefined','email=undefined; password=undefined','F; M1'),
('Argumentos null','email=null; password=null','F; M1'),
('Correo con espacios exteriores','email=" paciente@example.test "; password=K3','V'),
('Contraseña de espacios','email=C3; password=3 espacios','V'),
]
for field in ['email','password','fullName','identityCard','phoneNumber','address']:
    data.append(('Campo común vacío',f'P0 con {field}=""','F; M4'))
data.extend([
('Paciente completo','P0','V'),('Médico sin especialidad','D0 con selectedSpecialty=""','F; M5'),
('Médico sin título','D0 con document=null','F; M5'),('Médico completo','D0','V'),
('Prioridad de campos comunes','D0 con email="" y document=null','F; M4'),
('Nombre de solo espacios','P0 con fullName=3 espacios','V'),
('Teléfono vacío','phoneNumber=""; address=DIR','F; M6'),
('Dirección vacía','phoneNumber=TEL; address=""','F; M6'),
('Campos completos','phoneNumber=TEL; address=DIR','V'),
('Campos de solo espacios','phoneNumber=3 espacios; address=3 espacios','V'),
('Nombre vacío','fullName=""; phoneNumber=TEL','F; M7'),
('Celular vacío','fullName=NOM; phoneNumber=""','F; M7'),
('Nombre de solo espacios','fullName=3 espacios; phoneNumber=TEL','F; M7'),
('Celular de solo espacios','fullName=NOM; phoneNumber=3 espacios','F; M7'),
('Campos con espacios exteriores','fullName=" Médico de Prueba "; phoneNumber=" 00000000 "','V')])
assert len(data)==30

def case_table(title,fn,start,end,legend,evidence):
    head(title)
    p('Función evaluada: '+fn+'. '+legend)
    caption('Tabla. Casos de '+title.lower())
    rows=[]
    for n in range(start,end+1):
        scenario,inputs,result=data[n-1]
        state='Aprobado en Jest' if n not in [21,25] else 'Aprobado como caracterización; limitación funcional pendiente'
        rows.append([f'PU-{n:02d}',scenario,inputs,'A' if n<=5 else 'B',result,result,state])
    table(['ID','Escenario','Datos de entrada ficticios','Pasos','Resultado esperado','Resultado obtenido','Estado'],rows,[1.4,2.7,3.55,1.4,2.15,2.15,3.24])
    source('elaboración propia a partir de '+evidence+', 2026. Los pasos A/B y resultados V/F/M se definen en las convenciones del apéndice.')

case_table('Inicio de sesión','validateLoginFields',1,9,
 'C1 = paciente@telemedicina.gob.bo; C2 = medico@telemedicina.gob.bo; C3 = paciente@example.test. K1 = Password123#; K2 = ClaveSegura2026!; K3 = Ficticia123!. Todas son entradas literales ficticias de los archivos de prueba. La contraseña no se recorta. PU-09 comprueba presencia de caracteres; no autentica esa contraseña.', 'E1, E2, E3 y E4')
case_table('Registro','validateRegistrationFields',10,21,
 'P0 es el objeto patient: email=paciente@example.test; password=Ficticia123!; fullName=Persona de Prueba; identityCard=TEST-001; phoneNumber=00000000; address=Dirección ficticia; isDoctor=false; selectedSpecialty=""; document=null. D0 conserva esos campos y establece isDoctor=true, selectedSpecialty=especialidad-ficticia y document={name: titulo-ficticio.pdf, uri: file:///ficticio.pdf}. Cada fila modifica únicamente los campos indicados.', 'E1, E2 y E4')
case_table('Perfil de paciente','validatePatientProfileFields',22,25,
 'TEL = "00000000" y DIR = "Dirección ficticia". La función recibe phoneNumber y address, en ese orden. PU-25 caracteriza H-02 y no demuestra que los datos sean adecuados para guardar un perfil.', 'E1, E2 y E4')
case_table('Perfil de médico','validateDoctorProfileFields',26,30,
 'TEL = "00000000" y NOM = "Médico de Prueba". La función recibe fullName y phoneNumber, en ese orden. La comprobación de presencia aplica trim() a ambos valores; no se prueba la persistencia de las modificaciones del perfil.', 'E1, E2 y E4')

d.add_page_break();head('Hallazgos funcionales pendientes')
head('H-01 Nombre de solo espacios en el registro',3)
p('PU-21 pasó como prueba de caracterización: al sustituir fullName por tres espacios en P0, validateRegistrationFields devolvió un objeto con isValid=true. La condición comprueba !fullName sin recortar espacios, por lo que una cadena no vacía supera la regla. El resultado evidencia una limitación de la validación local; no demuestra que se haya creado una cuenta con esos datos. Estado: pendiente de corrección y de una nueva verificación. Evidencias: E1, E2, E4 y E6.')
head('H-02 Teléfono y dirección de solo espacios en el perfil de paciente',3)
p('PU-25 pasó como prueba de caracterización: validatePatientProfileFields recibió tres espacios en cada argumento y devolvió un objeto con isValid=true. Las condiciones !phoneNumber y !address no detectan cadenas compuestas únicamente por espacios. No se ejecutó una actualización de perfil en Supabase. Estado: pendiente de corrección y de una nueva verificación. Evidencias: E1, E2, E4 y E6.')
p('La eventual corrección requiere definir la regla de presencia de estos campos, implementarla y volver a ejecutar las pruebas correspondientes.')
head('Alcance de las conclusiones')
p('Los 30 resultados confirman el comportamiento de las cuatro funciones únicamente frente a las entradas evaluadas. No permiten concluir que todo el módulo de usuarios funcione correctamente ni que el sistema esté libre de errores. No se evaluaron autenticación real, roles, duplicados de usuarios, carga de títulos, persistencia, interfaz ni otros módulos. Los dos casos de caracterización mantienen sus limitaciones funcionales pendientes aunque Jest los haya marcado como aprobados.')
head('Evidencia gráfica de los casos y aserciones')
figure_slot('Figura. Casos y aserciones de la validación de inicio de sesión','La evidencia corresponde a loginValidation.test.ts y al bloque de límites del login de userValidation.test.ts. Permite relacionar PU-01 a PU-09 con la función invocada y las aserciones.')
figure_slot('Figura. Casos y aserciones de registro y perfiles','La evidencia corresponde a los bloques de registro, perfil de paciente y perfil médico de userValidation.test.ts. PU-21 y PU-25 identifican las dos caracterizaciones. La salida global de Jest se referencia en el apartado 3.3.1, sin duplicar la figura de ejecución.')
OUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUT)
audit={'counts':{'login':9,'registro':12,'perfil_paciente':4,'perfil_medico':5},'case_titles':cases,
 'source_hashes':{f:hashlib.sha256((ROOT/f).read_bytes()).hexdigest() for f in sources},
 'reference_hash':hashlib.sha256(Path(r'C:\Users\inutil\Downloads\Documento_Oficial_V2 (5).docx').read_bytes()).hexdigest(),
 'missing_captures':True,'appendix_number':None}
(ROOT/'tmp/word-usuarios/verificacion.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT)
