from catalog import *
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=ROOT/'output/documents';OUT.mkdir(parents=True,exist_ok=True)
BASE=Document(TEMPLATE)
WIDTH=17.272
FIVE=[1.9,3.1,4.9,4.3,3.072]
SOURCE='Fuente: Elaboración propia, 2026'
levels=[0]*4;section=None
for para in Document(OFFICIAL).paragraphs:
 if not para.style.name.startswith('Heading '):continue
 n=int(para.style.name.split()[-1])-1
 if n>3:continue
 x=para._p.find('w:pPr/w:numPr/w:numId',para._p.nsmap)
 if x is not None and x.get(qn('w:val'))=='0':continue
 levels[n]+=1
 for j in range(n+1,4):levels[j]=0
 if levels[0]==3 and para.text.strip()=='PRUEBAS UNITARIAS':section='.'.join(map(str,levels[:n+1]))
assert section=='3.3.1'

def start(filename,title):
 path=OUT/filename;shutil.copyfile(TEMPLATE,path);d=Document(path)
 for el in list(d._element.body):
  if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
 for name,size in [('Normal',11),('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',11.5),('Caption',10.5)]:
  st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.05
  st.paragraph_format.keep_with_next=name.startswith('Heading') or name=='Caption'
  for bd in list(st.element.iter(qn('w:pBdr'))):bd.getparent().remove(bd)
 d.styles['Normal'].paragraph_format.keep_together=False
 d.core_properties.title=title;d.core_properties.subject='Pruebas unitarias del sistema de telemedicina'
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments=''
 return d,path

def p(d,text='',style=None,bold=False,center=False):
 q=d.add_paragraph(text,style);q.paragraph_format.widow_control=True
 q.paragraph_format.keep_together=True
 if center:q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if bold:
  q.paragraph_format.keep_with_next=True
  for run in q.runs:run.bold=True
 return q
def h(d,text,level=2,new=False):
 q=p(d,text,'Heading '+str(level));q.paragraph_format.page_break_before=new;return q
def caption(d,text):return p(d,text,'Caption',bold=True,center=True)
def source(d):
 q=p(d,SOURCE,center=True);q.paragraph_format.space_before=Pt(4);q.paragraph_format.space_after=Pt(9)
 for r in q.runs:r.font.size=Pt(9.5);r.italic=True
 return q
def wrap_paths(text):
 # Word may break long paths; the original command remains in evidence metadata.
 return str(text).replace('/','/\u200b').replace('\\','\\\u200b')

def table(d,headers,rows,widths,header=True,font=10.5):
 t=d.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 # Use the source's table component with normalized width.
 original=t._tbl.tblPr;t._tbl.remove(original);t._tbl.insert(0,deepcopy(BASE.tables[0]._tbl.tblPr))
 pr=t._tbl.tblPr
 for tag in ['w:tblInd','w:tblW','w:tblBorders','w:tblCellMar','w:tblLayout']:
  for el in list(pr.findall(tag,pr.nsmap)):pr.remove(el)
 tw=OxmlElement('w:tblW');tw.set(qn('w:w'),str(round(WIDTH*567)));tw.set(qn('w:type'),'dxa');pr.append(tw)
 lay=OxmlElement('w:tblLayout');lay.set(qn('w:type'),'fixed');pr.append(lay)
 bd=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'B7B7B7');bd.append(z)
 pr.append(bd)
 mar=OxmlElement('w:tblCellMar')
 for name,val in [('top',65),('bottom',65),('left',80),('right',80)]:
  el=OxmlElement('w:'+name);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
 pr.append(mar)
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 allrows=([headers] if header else [])+rows
 for i,values in enumerate(allrows):
  row=t.add_row();trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
  if i==0:
   hd=OxmlElement('w:tblHeader');trpr.append(hd)
  for j,(cell,value,w) in enumerate(zip(row.cells,values,widths)):
   cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cell.text=wrap_paths(value)
   for q in cell.paragraphs:
    q.paragraph_format.space_before=Pt(0);q.paragraph_format.space_after=Pt(2);q.paragraph_format.line_spacing=1.0
    q.paragraph_format.keep_with_next=(i==0);q.paragraph_format.widow_control=True
    for r in q.runs:r.font.name='Arial';r.font.size=Pt(font);r.bold=(i==0)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8B8B8');cell._tc.get_or_add_tcPr().append(sh)
 for cell in t.rows[-1].cells:
  for q in cell.paragraphs:q.paragraph_format.keep_with_next=True
 return t

def cover(d,title,subtitle):
 p(d,'UNIVERSIDAD PRIVADA DEL VALLE',bold=True,center=True)
 p(d,'FACULTAD DE INFORMÁTICA Y ELECTRÓNICA',bold=True,center=True)
 p(d,'CARRERA DE INGENIERÍA DE SISTEMAS INFORMÁTICOS',bold=True,center=True)
 q=p(d,title,'Title',center=True);q.paragraph_format.space_before=Pt(32);q.paragraph_format.space_after=Pt(16)
 p(d,subtitle,center=True)
 q=p(d,'Sistema de telemedicina con triaje mediante modelos de inteligencia artificial',center=True);q.paragraph_format.space_before=Pt(34)
 p(d,'Cochabamba - Bolivia',center=True);p(d,'2026',center=True)

v=executions['usuarios']['versiones']
toolsrows=[['Jest',v['jest'],'Ejecuta las pruebas de las funciones JavaScript y TypeScript.'],['jest-expo',v['jest-expo'],'Adapta Jest al entorno de la aplicación Expo.'],['Python y unittest','3.12.14','Ejecutan las pruebas aisladas del servicio de análisis; unittest forma parte de Python.'],['Node.js y npm',v['node'].lstrip('v')+' / '+v['npm'],'Ejecutan las herramientas y administran las dependencias.'],['TypeScript',v['typescript'],'Permite comprobar los tipos del proyecto.'],['Expo',v['expo'],'Es una tecnología de la aplicación; no es el ejecutor de pruebas.']]

figures=[dict(id='3.1',key='config',title='Configuración de Jest utilizada en las pruebas',description='La configuración selecciona jest-expo y los patrones de archivos de prueba.',file='jest.config.js',visible='Archivo completo: preset, testMatch y transformIgnorePatterns.',demonstrates='Configuración real del proyecto.'),
 dict(id='3.2',key='usuarios',title='Casos de validación de usuarios y comprobación de sus respuestas',description='Los casos comparan las respuestas locales de los validadores con las expectativas definidas.',file='__tests__/userValidation.test.ts',visible='Importaciones y bloque PU-21; una segunda toma del bloque PU-25 si no cabe a tamaño legible.',demonstrates='Pruebas de caracterización que conservan las expectativas originales.')]
for index,(key,m) in enumerate(modules.items(),3):
 figures.append(dict(id=f'3.{index}',key=key,title='Ejecución de pruebas de '+('triaje en Python' if key=='triaje-python' else m['name'].lower()),description=f"La ejecución registrada contiene {m['count']} casos aprobados. "+('Se utilizó unittest con el proveedor sustituido.' if key=='triaje-python' else 'Se utilizó Jest sobre los archivos seleccionados.'),file='Terminal PowerShell',visible='Comando completo, nombres de los archivos o casos y resumen final; fecha de la ejecución visible.',demonstrates=f"Resultado real de {m['count']} casos; no certifica el módulo completo."))
for index,key in enumerate([k for k in modules if k!='usuarios'],1):
 m=modules[key];figures.append(dict(id=f'A.{index}',key=key,title='Código de pruebas de '+('triaje en Python' if key=='triaje-python' else m['name'].lower()),description='Las entradas ficticias, la llamada a la función y las comprobaciones documentan el alcance de la evaluación.',file=m['tests'][0],visible={'triaje':'PU-31 a PU-37, en dos tomas si es necesario.','triaje-python':'setUp y sustitución de red; PU-40 y PU-41.','videoconsultas':'PU-45 a PU-51.','historial-recetas':'PU-55 a PU-58; PU-59 y PU-60 en otra toma.','pagos':'jest.mock de Supabase, PU-62 a PU-64.','reportes':'PU-65 a PU-69.'}[key],demonstrates='Lógica evaluada y aserciones; datos y servicios externos ficticios.'))

guide=['# Guía de capturas reales','', 'Las capturas están pendientes. Completar una etapa por vez. Mantener visibles el nombre del archivo o el comando, el contenido y su tamaño legible. No abrir .env ni mostrar credenciales reales. Una captura de código o de una ejecución nueva no debe presentarse como una imagen tomada en la ejecución histórica del 30 de septiembre. Los resultados originales se conservan en evidencias/pruebas-completas.','',f'Carpeta del proyecto: `{ROOT}`.','', 'Para cada toma: abrir el archivo en el editor, ajustar el tamaño de texto y usar Windows + Mayús + S. Guardar la imagen sin recortar nombres o resultados importantes. Las figuras de ejecución pueden dividirse en tomas contiguas, sin alterar los resultados.','']
for key in modules:(E/key/'capturas.md').write_text('# Capturas pendientes\n\n',encoding='utf-8')
for f in figures:
 g=[f"## Figura {f['id']} {f['title']}",'',f"Abrir: `{str(ROOT/f['file']) if f['file']!='Terminal PowerShell' else f['file']}`.",f"Visible: {f['visible']}",f"Demuestra: {f['demonstrates']}",SOURCE+'.','']
 if f['file']=='Terminal PowerShell':
  key=f['key'];m=modules[key]
  if key=='triaje-python':command='& "C:\\Users\\inutil\\.cache\\codex-runtimes\\codex-primary-runtime\\dependencies\\python\\python.exe" -m unittest discover -s ai-engine -p test_unit_llm.py -v'
  else:command='node node_modules/jest/bin/jest.js --runInBand --watch=false --verbose --runTestsByPath '+' '.join(m['tests'])
  g+=['Desde la carpeta del proyecto, ejecutar estos comandos. Esta nueva ejecución no sobrescribe el JSON histórico:','```powershell','Get-Date -Format o',command,'```','Registrar la fecha de la nueva captura y conservar su salida antes de incorporarla. Los 69 resultados de los documentos corresponden a las ejecuciones identificadas en ejecucion.json.','']
 guide+=g
 if f['key'] in modules:
  cp=E/f['key']/'capturas.md'
  with cp.open('a' if cp.exists() else 'w',encoding='utf-8') as fp:fp.write('\n'.join(g)+'\n')
(OUT/'Guia_capturas_pruebas_unitarias.md').write_text('\n'.join(guide),encoding='utf-8')
save(E/'figuras-pendientes.json',figures)

def figure(d,f):
 caption(d,f"Figura {f['id']}. {f['title']}")
 q=p(d,'[Espacio reservado para captura real pendiente de incorporación]',center=True)
 q.paragraph_format.space_before=Pt(14);q.paragraph_format.space_after=Pt(14);q.paragraph_format.keep_with_next=True
 for r in q.runs:r.italic=True;r.font.size=Pt(10.5)
 q=p(d,f['description']);q.paragraph_format.keep_with_next=True;source(d)

def findings_text(d):
 for ident,ids,text in findings:
  p(d,f"{ident}. "+text+' Estado: pendiente de corrección y nueva comprobación. '+', '.join(f'PU-{n:02}' for n in ids)+'.')

pending=[
 ('Asignar roles y especialidades','Requiere integración','La asignación real, autorización y persistencia de roles y especialidades no se ejecutaron. Las validaciones del registro no acreditan estos cambios.'),
 ('Controlar disponibilidad médica','Pendiente','Existe el control de disponibilidad. El fallo manual informado —el botón no cambia el estado— no se reprodujo en esta ejecución ni se declara corregido. Se requiere comprobar interfaz, sesión y actualización en Supabase.'),
 ('Transmitir videoconsulta','Requiere integración','La aplicación utiliza el componente de videollamada, pero no se estableció una conexión entre médico y paciente ni se comprobaron audio y video.'),
 ('Generar código QR de pago','No implementado','La pantalla muestra un icono QR y una simulación. No se identificó generación de un código bancario verificable.'),
]

# Documento académico RUP 2.
d,path=start('Pruebas_unitarias_RUP2_resultados_actualizados.docx','Pruebas unitarias del sistema de telemedicina')
cover(d,'Pruebas unitarias del sistema de telemedicina','Contenido para el Capítulo III y el apéndice de pruebas')
h(d,'CAPÍTULO III',1,True);h(d,'PRUEBAS DE CALIDAD',1);h(d,section+' PRUEBAS UNITARIAS',2)
p(d,'Se evaluaron funciones aisladas de las validaciones de usuarios, la preparación del triaje, las decisiones de la sala de espera, el historial y las recetas, la simulación de pagos y los cálculos de reportes. Las pruebas utilizaron datos ficticios y sustitutos de los servicios externos para comparar las respuestas del código con las expectativas de cada caso.')
p(d,'Las ejecuciones finales del 30 de septiembre de 2026 registraron 69 casos aprobados: 62 en siete suites de Jest y 7 en unittest de Python. No hubo casos fallidos ni omitidos en esas ejecuciones finales. La aprobación indica coincidencia con las aserciones; no demuestra que un módulo completo funcione correctamente.')
caption(d,'Tabla 3.1. Herramientas y librerías utilizadas en las pruebas unitarias');table(d,['Herramienta o tecnología','Versión verificada','Función'],toolsrows,[3.4,2.75,11.122]);source(d)
caption(d,'Tabla 3.2. Resumen de pruebas unitarias de las funciones evaluadas')
table(d,None,[['Objetivo','Comprobar de forma aislada las reglas implementadas y registrar las limitaciones reproducidas por los casos de prueba.'],['Acción','Ejecutar pruebas sobre gestión de usuarios (30 casos), triaje (14), videoconsultas (7), historial y recetas (9), simulación de pagos (4) y reportes (5). Los servicios de autenticación, datos, inteligencia artificial y pago no se utilizaron en producción.'],['Efecto','Los 69 casos coincidieron con sus expectativas. Se conservaron los hallazgos del registro y del perfil del paciente y se documentaron otras limitaciones de las reglas actuales. Las operaciones reales de disponibilidad, videollamada, persistencia y pago no quedan aprobadas por estos resultados. El detalle se presenta en el apéndice de pruebas unitarias.']],[2.8,14.472],header=False,font=11);source(d)
h(d,'Evidencias gráficas de las pruebas unitarias',2)
p(d,'Las siguientes figuras corresponden a la configuración, los casos y las ejecuciones por grupo. Las imágenes están pendientes de incorporación; los resultados se respaldan en los registros de ejecución conservados.')
for f in figures:
 if f['id'].startswith('3.'):figure(d,f)
p(d,'Los escenarios, pasos y resultados esperados se detallan en el apéndice siguiente. Los casos que requieren servicios, interfaz o evaluación clínica se distinguen de las pruebas unitarias ejecutadas.')
h(d,'APÉNDICE',1,True);h(d,'PRUEBAS UNITARIAS DEL SISTEMA DE TELEMEDICINA',2)
p(d,'Cada fila corresponde a una prueba automatizada de una función. Jest o unittest prepararon las entradas, invocaron el código y compararon la respuesta con la expectativa. No se realizaron recorridos manuales de pantalla. «Éxito» indica que la aserción se cumplió; los casos de caracterización pueden tener éxito al reproducir una limitación funcional.')
p(d,'Se conservan los identificadores PU-01 a PU-30 de las validaciones de usuarios. Los casos incorporados se identifican desde PU-31 hasta PU-69. Las actividades con estado Pendiente, Requiere integración o No implementado no se suman a esos 69 casos ni se consideran omitidas por Jest.')
order=['Iniciar sesión','Registrar usuario','Asignar roles y especialidades','Editar perfil del paciente','Editar perfil del médico',capture,process,classify,waiting,'Controlar disponibilidad médica','Transmitir videoconsulta',pred,clinical,recipe,'Generar código QR de pago',pay,report]
notes={
'Iniciar sesión':'Se evaluó la validación de campos obligatorios; no se autenticó una cuenta en Supabase.',
'Registrar usuario':'Se evaluaron campos comunes y requisitos locales del médico; no se crearon cuentas ni se cargaron títulos.',
'Editar perfil del paciente':'Se evaluaron teléfono y dirección antes de guardar; no se modificó información persistida.',
'Editar perfil del médico':'Se evaluaron nombre y celular antes de guardar; no se modificó información persistida.',
capture:'Se evaluaron la limpieza de espacios y la composición del relato. La captura de voz y la interfaz requieren integración.',
process:'Se utilizó unittest para comprobar el procesamiento de respuestas simuladas. No se evaluó traducción ni transcripción de un proveedor real.',
classify:'Se comprobaron el filtro local de especialidad y dos reglas del servicio Python. No se midió exactitud clínica ni disponibilidad real de especialistas.',
waiting:'Se probaron decisiones locales de cierre y habilitación. La consulta de la cola, sus cambios en tiempo real y la posición mostrada requieren integración.',
pred:'Se evaluó la preparación de los datos para su presentación; no se comprobó su visualización ni su validez clínica.',
clinical:'Se evaluaron únicamente diagnóstico y tratamiento obligatorios. Guardar y recuperar el historial requiere integración.',
recipe:'Se evaluó el contenido HTML utilizado para preparar la receta. La generación, visualización y distribución del PDF requieren integración y revisión manual.',
pay:'La verificación bancaria real no está implementada. Los cuatro casos siguientes evalúan únicamente la simulación existente con Supabase sustituido.',
report:'El módulo tiene una implementación parcial. Se probaron sus cálculos locales; la consulta real requiere integración. Los filtros por fecha y la exportación de reportes no están implementados.'}
tn=0
for use in order:
 h(d,'Caso de uso '+use,3)
 cs=[c for c in cases if c['use']==use]
 if cs:
  p(d,notes[use]);tn+=1;caption(d,f'Tabla A.{tn}. Casos de {use[0].lower()+use[1:]}')
  table(d,['Caso de prueba','Escenario de prueba','Pasos de la prueba','Resultado esperado','Éxito / Fracaso'],[[c['id'],c['scenario'],c['steps'],c['expected'],c['state']] for c in cs],FIVE)
  source(d)
  folders=list(dict.fromkeys(c['module'] for c in cs));p(d,'Evidencia: '+', '.join('evidencias/pruebas-completas/'+f for f in folders)+'. Se conservaron resultado.json, salida.log, ejecucion.json y las fuentes de las pruebas.')
 else:
  _,state,note=next(v for v in pending if v[0]==use);p(d,state+'. '+note)
 if use=='Editar perfil del médico':p(d,'Los casos de usuarios se relacionan con las figuras 3.2 y 3.3. PU-21 y PU-25 reproducen limitaciones pendientes; sus resultados no acreditan la calidad de los datos aceptados.')
 # Place code evidence at the end of its related use-case group.
 figkey={capture:'triaje',process:'triaje-python',waiting:'videoconsultas',recipe:'historial-recetas',pay:'pagos',report:'reportes'}.get(use)
 if figkey:
  f=next(f for f in figures if f['key']==figkey and f['id'].startswith('A.'));figure(d,f)
h(d,'Hallazgos y alcance de los resultados',2)
findings_text(d)
p(d,'Además, el fallo manual informado del botón de disponibilidad permanece pendiente de reproducción y corrección. La evaluación no acredita autenticación, persistencia, pago bancario, transmisión de videollamada ni precisión médica de las respuestas de inteligencia artificial. Tampoco se midió cobertura de líneas o ramas, por lo que no se afirma una cobertura completa del código.')
# Usar palabras completas en el extracto académico; el informe conserva los nombres del código.
plain={'standardized_symptoms':'síntomas estandarizados','urgency_level':'nivel de urgencia','recommended_specialty':'especialidad recomendada','completedTriages':'atenciones completadas','totalRevenue':'ingreso estimado'}
for element in d._element.body.iter(qn('w:t')):
 if element.text:
  for term,label in plain.items():element.text=element.text.replace(term,label)
d.save(path)
rup_path=path

# Informe técnico con tablas transpuestas para mantener una tipografía legible.
d,path=start('Informe_detallado_pruebas_unitarias_telemedicina.docx','Informe detallado de pruebas unitarias de telemedicina')
cover(d,'Informe detallado de pruebas unitarias de telemedicina','Resultados, alcance, cambios y evidencias de ejecución')
h(d,'Resultados de la ejecución',1,True)
p(d,'Las ejecuciones finales verificaron 69 casos únicos: 62 en siete suites de Jest y 7 en unittest de Python. Todos cumplieron sus aserciones, sin fallidos ni omitidos en el resultado final. Este resultado comprende validaciones, decisiones locales, preparación de contenido y servicios sustituidos; no representa la aprobación de módulos completos.')
p(d,'La primera ejecución de Python registró siete errores de preparación por la resolución del módulo dotenv al instalar los sustitutos. Se conservó ese intento, se corrigió únicamente el mecanismo de sustitución y se repitieron los siete casos, que aprobaron. Los siete intentos iniciales no se suman como nuevos casos.')
caption(d,'Tabla 1. Resultados finales por grupo de ejecución')
table(d,['Grupo','Ejecutor','Casos aprobados','Fallidos / omitidos'],[[('Triaje del servicio Python' if k=='triaje-python' else m['name']),('unittest' if k=='triaje-python' else 'Jest'),str(m['count']),'0 / 0'] for k,m in modules.items()],[6,3,3.5,4.772]);source(d)
h(d,'Herramientas y configuración',2)
caption(d,'Tabla 2. Versiones instaladas verificadas');table(d,['Herramienta o tecnología','Versión','Uso'],toolsrows,[3.4,2.75,11.122]);source(d)
p(d,'Jest utiliza el preset jest-expo, los patrones de testMatch y las extensiones ts, tsx, js, jsx y json de jest.config.js. Las ejecuciones seleccionaron archivos explícitos con --runTestsByPath, desactivaron el modo de observación y se ejecutaron en un solo proceso. No se ejecutaron indiscriminadamente otros scripts del repositorio.')
p(d,'Las pruebas de pagos sustituyen el módulo lib/supabase antes de importarlo. Las pruebas Python sustituyen dotenv.load_dotenv, urllib.request.urlopen, os.getenv y la escritura de errores. No se cargó el archivo de variables de entorno ni se realizaron solicitudes reales al proveedor. El servicio Python se probó sin iniciar FastAPI.')
p(d,'La comprobación estática con tsc --noEmit --pretty false finalizó con código 0, sin diagnósticos. git diff --check no detectó errores de espacios; sus avisos sobre LF/CRLF no son fallos de las pruebas. Estas comprobaciones complementarias no se suman a los 69 casos.')
h(d,'Identificación de las ejecuciones',1)
p(d,'Fecha: 30 de septiembre de 2026. Entorno: Windows, PowerShell y zona America/La_Paz, UTC−04:00. Los tiempos de la tabla corresponden al inicio local de cada ejecución final; los registros JSON conservan inicio y fin en UTC.')
p(d,'Referencia del repositorio: '+executions['usuarios']['commit']+'. Había cambios locales antes de ampliar las pruebas, por lo que el identificador del commit no representa por sí solo la versión ejecutada. Se conservaron inventarios por ejecución, fuentes copiadas y hashes de respaldo.')
caption(d,'Tabla 3. Archivos y fecha de las ejecuciones finales')
rows=[]
for k,m in modules.items():
 local=datetime.fromisoformat(executions[k]['inicio'].replace('Z','+00:00')).astimezone(timezone(timedelta(hours=-4)))
 rows.append([k,local.strftime('%d/%m/%Y\n%H:%M:%S'), '\n'.join(m['tests'])])
table(d,['Carpeta de evidencia','Inicio UTC−04:00','Archivos ejecutados'],rows,[4,3,10.272]);source(d)
p(d,'Carpeta raíz de evidencia: evidencias/pruebas-completas. Cada subcarpeta contiene ejecucion.json con comando, versiones, fecha y referencia local; resultado.json con estados; salida.log con la salida textual; fuentes con las copias utilizadas; archivos.json, alcance.md y capturas.md. La carpeta triaje-python conserva también intento-inicial.json e intento-inicial.log.')
h(d,'Comandos registrados',2)
for k,m in modules.items():
 p(d,k,bold=True)
 q=p(d,wrap_paths(executions[k]['comando']))
 for run in q.runs:run.font.size=Pt(10)
p(d,'Los comandos anteriores documentan la ejecución ya realizada. La guía de capturas propone ejecuciones adicionales que no sobrescriben esos resultados históricos.')
h(d,'Módulos y casos de uso revisados',1)
for key,m in modules.items():
 if key=='triaje-python':continue
 h(d,m['name'],2);p(d,m['limits'])
 p(d,'Funciones probadas: '+', '.join(dict.fromkeys(c['fn'] for c in cases if c['module']==key or (key=='triaje' and c['module']=='triaje-python')))+'.')
 p(d,'Casos de uso con alcance unitario: '+', '.join(dict.fromkeys(c['use'] for c in cases if c['module']==key or (key=='triaje' and c['module']=='triaje-python')))+'.')
h(d,'Cambios y repetición de pruebas',1)
p(d,'Para permitir pruebas aisladas, se extrajeron reglas existentes de sus componentes y se reemplazaron los bloques originales por llamadas a esas funciones. Se conservaron sus condiciones y resultados actuales, incluidas las limitaciones. No se corrigieron H-01 ni H-02 ni se cambiaron expectativas para obtener resultados aprobados.')
caption(d,'Tabla 4. Archivos creados y pantallas adaptadas en esta ampliación')
rows=[]
for k,m in modules.items():
 if k=='usuarios':continue
 new=m['tests']+(m['sources'] if k!='triaje-python' else [])
 rows.append([k,'\n'.join(new),'\n'.join(m['modified']) or 'Ninguna; llm_service.py se conserva.'])
table(d,['Grupo','Archivos nuevos','Archivos modificados'],rows,[3.4,6.6,7.272]);source(d)
p(d,'Se creó scripts/record-unit-tests.cjs para ejecutar archivos seleccionados y guardar resultados. Se crearon también los archivos de evidencia y los dos documentos independientes. Las fuentes de usuarios y sus dos archivos de prueba ya existían al comenzar esta ampliación y se volvieron a ejecutar sin cambiar sus expectativas. Las modificaciones previas en registro, perfiles, package.json, package-lock.json y tsconfig.json no se atribuyen a esta ampliación.')
h(d,'Error de preparación y corrección mínima',2)
p(d,'El intento inicial del servicio Python tuvo siete errores antes de ejecutar las aserciones: ModuleNotFoundError para dotenv. La restauración de sys.modules al terminar la sustitución de dotenv provocó que patch con una ruta textual intentara importar llm_service de nuevo. Se cambió a patch.object sobre el módulo ya cargado, manteniendo sustituidas la red y la lectura de configuración. No se modificó llm_service.py ni las expectativas de los casos.')
p(d,'La repetición final de los siete casos Python aprobó sin errores. La evidencia del fallo y de la repetición se conserva en triaje-python. La ejecución de videoconsultas quedó registrada entre el intento inicial de Python y su repetición final; los tiempos de cada registro permiten reconstruir esa secuencia.')
h(d,'Hallazgos funcionales pendientes',1);findings_text(d)
h(d,'Actividades que no son pruebas unitarias ejecutadas',2)
for use,state,note in pending:p(d,use+'. '+state+'. '+note)
p(d,'Autenticar usuarios, crear cuentas, guardar perfiles y comprobar permisos requieren integración. En triaje, la calidad de la traducción, la clasificación de urgencia y la derivación necesitan casos clínicos y revisión profesional; las respuestas simuladas no son una validación médica. El procesamiento real de voz requiere integración.')
p(d,'Guardar y recuperar el historial, generar y compartir un PDF real requieren integración y revisión manual. En pagos, la verificación bancaria está sin implementar. En reportes, falta implementar filtros por fecha y exportación; la consulta real de datos permanece sin comprobar. No se midió cobertura de líneas, ramas o seguridad del sistema.')
h(d,'Detalle de los casos ejecutados',1,True)
p(d,'Las tablas presentan los diez campos de cada caso en orientación vertical, agrupando tres identificadores por tabla para conservar la legibilidad. «Resultado obtenido» expresa las comprobaciones respaldadas por las aserciones aprobadas; no se registraron objetos adicionales no comprobados por los tests. Las rutas de evidencia se resuelven dentro de evidencias/pruebas-completas.')
p(d,'Los registros de paciente usados en usuarios tienen correo paciente@example.test, contraseña ficticia Ficticia123!, nombre Persona de Prueba, identidad TEST-001, teléfono 00000000 y Dirección ficticia. Los de médico agregan especialidad-ficticia y un objeto de documento titulo-ficticio.pdf que no se abre ni se carga. Los literales de los cinco casos originales también son datos de prueba, no credenciales verificadas.')
for offset in range(0,69,3):
 batch=cases[offset:offset+3]
 if offset>0:
  q=caption(d,f"Tabla {5+offset//3}. Detalle de {batch[0]['id']} a {batch[-1]['id']}");q.paragraph_format.page_break_before=True
 else:caption(d,'Tabla 5. Detalle de PU-01 a PU-03')
 rows=[]
 for label,field in [('Módulo','module_name'),('Caso de uso','use'),('Función','fn'),('Escenario','scenario'),('Datos de prueba','data'),('Resultado esperado','expected'),('Resultado obtenido','actual'),('Estado','state'),('Evidencia','evidence')]:
  rows.append([label]+[c[field] for c in batch])
 table(d,['Identificador']+[c['id'] for c in batch],rows,[2.8,4.824,4.824,4.824],font=10.5);source(d)
h(d,'Respaldo documental y evidencias gráficas',1,True)
p(d,'Documento oficial vigente: Documento_Oficial_V2 (5).docx, que sustituye al archivo anterior. La jerarquía existente ubica Pruebas unitarias en 3.3.1. No se encontró una secuencia de apéndices en el índice vigente, por lo que el extracto utiliza el título APÉNDICE sin copiar el número 3 del modelo. La numeración de tablas y figuras identifica el extracto independiente.')
p(d,'Referencias de presentación: PG - MODLEO_UML_RUP 2 1 (1).pdf, página impresa 147 para Objetivo, Acción y Efecto y página impresa 184 para las cinco columnas de casos; plantilla Demo_pruebas_unitarias_RUP2_documento_completo.docx. Los documentos de referencia permanecen sin sobrescribir.')
p(d,'Las capturas reales están pendientes. Sus espacios se identifican en Pruebas_unitarias_RUP2_resultados_actualizados.docx. La guía independiente Guia_capturas_pruebas_unitarias.md especifica el archivo, la parte visible, el propósito, el título y la fuente de cada figura. Los logs y JSON respaldan los resultados, pero no sustituyen las imágenes exigidas.')
p(d,'Los resultados finales muestran que las respuestas evaluadas coincidieron con las expectativas definidas. Los hallazgos de caracterización siguen pendientes, y las actividades de integración, las pruebas manuales y la evaluación clínica delimitan el trabajo aún necesario. No se concluye que el sistema esté libre de errores.')
d.save(path)

# Verificaciones estructurales y de preservación de referencias.
rd=Document(rup_path);body=rd._element.body;chaptertables=0;appendix=False
for el in body:
 if el.tag==qn('w:p') and ''.join(el.itertext())=='APÉNDICE':appendix=True
 if el.tag==qn('w:tbl') and not appendix:chaptertables+=1
# Use paragraph content rather than XML itertext duplication when identifying the boundary.
chaptertables=0
for el in body:
 textval=''.join(t.text or '' for t in el.iter(qn('w:t')))
 if textval=='APÉNDICE':break
 if el.tag==qn('w:tbl'):chaptertables+=1
assert chaptertables==2
caseids=[row.cells[0].text for t in rd.tables for row in t.rows if re.fullmatch(r'PU-\d{2}',row.cells[0].text)]
assert len(caseids)==69 and len(set(caseids))==69
technical=Document(path)
techids=[c.text for t in technical.tables for c in t.rows[0].cells if re.fullmatch(r'PU-\d{2}',c.text)]
assert len(techids)==69 and set(techids)==set(caseids)
refs=read(E/'referencias.json')
assert all(sha(Path(r['archivo']))==r['sha256'] for r in refs)
assert all((ROOT/f).read_bytes()==(ROOT/'evidencias/pruebas-usuarios/fuentes'/(f+'.txt')).read_bytes() for f in modules['usuarios']['tests']+modules['usuarios']['sources'])
save(TMP/'auditoria-documentos.json',dict(casosRUP=len(caseids),casosInforme=len(techids),tablasCapitulo=chaptertables,numeracion=section,referenciasPreservadas=True,capturasPendientes=len(figures),archivos=[str(rup_path),str(path)]))
print('Dos documentos creados; 69 casos en ambos; dos tablas en el capítulo; 15 capturas pendientes.')
