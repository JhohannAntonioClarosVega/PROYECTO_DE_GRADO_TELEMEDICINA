import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
E = ROOT / 'evidencias/pruebas-completas'
T = ROOT / 'tmp/pruebas-completas'
O = ROOT / 'output/documents'
read = lambda p: json.loads(p.read_text(encoding='utf-8-sig'))
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
summary = read(E / 'resumen.json')
assert summary['casos'] == 69 and summary['jest'] == 62 and summary['python'] == 7
totals = []
for group in ['usuarios', 'triaje', 'videoconsultas', 'historial-recetas', 'pagos', 'reportes']:
    result = read(E / group / 'resultado.json')
    assert result['numFailedTests'] == 0 and result['numPendingTests'] == 0
    assert result['numPassedTests'] == result['numTotalTests']
    totals.append(result['numTotalTests'])
assert totals == [30, 7, 7, 9, 4, 5]
result = read(E / 'triaje-python/resultado.json')
assert result['total'] == 7 and not any(result[k] for k in ['failures', 'errors', 'skipped'])
assert all(c['status'] == 'passed' for c in result['cases'])
for reference in read(E / 'referencias.json'):
    assert sha(Path(reference['archivo'])) == reference['sha256']
audit = read(T / 'auditoria-documentos.json')
assert audit['casosRUP'] == audit['casosInforme'] == 69
assert audit['tablasCapitulo'] == 2
artifacts = []
for filename, pdfname, expected_pages in [
    ('Pruebas_unitarias_RUP2_resultados_actualizados.docx', 'rup.pdf', 23),
    ('Informe_detallado_pruebas_unitarias_telemedicina.docx', 'informe.pdf', 31),
]:
    path = O / filename
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    texts += [c.text for table in doc.tables for row in table.rows for c in row.cells]
    ids = set(re.findall(r'PU-\d{2}', '\n'.join(texts)))
    assert ids == {f'PU-{n:02d}' for n in range(1, 70)}
    pages = len(PdfReader(str(T / pdfname)).pages)
    assert pages == expected_pages
    artifacts.append({'archivo': filename, 'sha256': sha(path), 'paginas': pages,
                      'revisionVisualTodasLasPaginas': True, 'textoCortadoDetectado': False,
                      'tablasIlegiblesDetectadas': False})
qa = {'fechaUTC': datetime.now(timezone.utc).isoformat(), 'archivos': artifacts,
      'metodo': 'Exportación de Microsoft Word a PDF y revisión de cada PNG de Poppler a 120 dpi.',
      'renderizadorEmpaquetado': 'No disponible por ausencia de LibreOffice; se utilizó Microsoft Word.',
      'casosCoincidentes': 69, 'tablasCapituloRUP': 2, 'referenciasOriginalesPreservadas': True,
      'figurasPendientes': 15, 'estado': 'Contenido revisado; evidencias gráficas pendientes de incorporación.'}
(E / 'revision-documentos.json').write_text(json.dumps(qa, ensure_ascii=False, indent=2), encoding='utf-8')
(E / 'README.md').write_text('''# Evidencias de pruebas unitarias

Ejecuciones finales del 30 de septiembre de 2026: 69 casos únicos aprobados; 62 en siete suites de Jest y siete en unittest de Python. No hubo fallidos ni omitidos en las ejecuciones finales.

| Grupo | Casos |
| --- | ---: |
| Usuarios | 30 |
| Triaje con Jest | 7 |
| Triaje con Python | 7 |
| Decisiones de sala de espera | 7 |
| Historial y recetas | 9 |
| Simulación de pagos | 4 |
| Cálculos de reportes | 5 |

Cada carpeta conserva resultado.json, salida.log, ejecucion.json, fuentes, archivos.json, alcance.md y capturas.md. El primer intento de Python produjo siete errores de preparación; se conserva como intento-inicial.json y su log. Se corrigió únicamente la preparación de sustitutos y se repitieron los siete casos sin cambiar las expectativas ni el servicio.

casos.json relaciona PU-01 a PU-69 con escenarios y evidencias. resumen.json consolida ejecuciones y límites. Los hallazgos H-01 a H-09 siguen pendientes de corrección y nueva comprobación. Aprobar una aserción de caracterización no acredita un comportamiento funcional adecuado. El fallo manual de disponibilidad continúa pendiente de reproducción y corrección.

No se acreditan autenticación real, persistencia, transmisión de video, transacciones bancarias, generación final de PDF ni exactitud clínica o lingüística. No se midió cobertura de líneas ni ramas. Los servicios externos fueron sustituidos.

verificacion-estatica.json registra TypeScript y git diff --check con salida cero. cambios-aplicacion.patch, antes/ y las copias de fuentes documentan la extracción de lógica existente para probarla aisladamente. El repositorio ya contenía cambios locales antes de la ampliación.

Los Word editables están en output/documents. revision-documentos.json registra sus hashes y revisión visual. Faltan las capturas reales correspondientes a 15 figuras; figuras-pendientes.json y output/documents/Guia_capturas_pruebas_unitarias.md indican cómo obtenerlas. Los logs no sustituyen esas imágenes. Los documentos aún no constituyen una entrega académica completa con evidencias gráficas.
''', encoding='utf-8')
print(json.dumps({'casos': 69, 'paginas': [a['paginas'] for a in artifacts], 'figurasPendientes': 15}, ensure_ascii=False))
