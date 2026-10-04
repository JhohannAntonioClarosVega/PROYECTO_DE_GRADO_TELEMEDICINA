from pathlib import Path
import json, hashlib, re, copy
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]
QA=Path(__file__).parent
SRC=ROOT/'output/documents/Pruebas_unitarias_RUP2_resultados_actualizados.docx'
OUT=SRC.with_name('Pruebas_unitarias_RUP2_documento_final.docx')
E=ROOT/'evidencias/revision-hallazgos-20260930/final'
read=lambda p: json.loads(p.read_text(encoding='utf-8-sig'))
sha=lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
original=sha(SRC)
j=read(E/'jest.json'); py=read(E/'python.json'); meta=read(E/'ejecucion.json')
assert (j['numTotalTests'],j['numPassedTests'],j['numFailedTests'],j['numPendingTests'])==(98,98,0,0)
assert (py['total'],py['passed'],py['failures'],py['errors'],py['skipped'])==(16,16,0,0,0)
assert all(sha(ROOT/x['archivo'])==x['sha256'] for x in meta['fuentes'])
assert '98 passed' in (E/'jest.log').read_text(encoding='utf-8-sig')
assert 'Ran 16 tests' in (E/'python.log').read_text(encoding='utf-8-sig')
d=Document(SRC); ps=list(d.paragraphs)
def txt(p,s):
    if p.runs:
        p.runs[0].text=s
        for r in p.runs[1:]: r.text=''
    else:p.add_run(s)
def cell(c,s):
    txt(c.paragraphs[0],s)
    for p in c.paragraphs[1:]:p._element.getparent().remove(p._element)
def setp(i,s):txt(ps[i],s)
def after(p,s,model=None):
    el=copy.deepcopy((model or p)._p);p._p.addnext(el);q=Paragraph(el,p._parent);txt(q,s);return q

setp(12,'La revisión final del 30 de septiembre de 2026 registró 114 pruebas ejecutadas y aprobadas: 98 en ocho suites de Jest y 16 en unittest de Python. No hubo pruebas fallidas ni omitidas. Estos resultados corresponden a las funciones y componentes aislados descritos en el apéndice; no acreditan el funcionamiento completo de los módulos.')
setp(29,'Se registraron 39 casos aprobados de validación local de usuarios en dos suites de Jest. No se comprobó la autenticación real.')
setp(31,'Figura 3.4. Ejecución de pruebas de preparación y filtro local del triaje con Jest')
setp(33,'Se registraron 7 casos aprobados de preparación y filtro del triaje. PU-37 conserva una caracterización cuya regla de asignación sigue pendiente.')
setp(37,'Se registraron 16 casos aprobados con unittest de Python y el proveedor sustituido. Se verificaron respuestas simuladas y errores; no se midió exactitud clínica.')
setp(39,'Figura 3.6. Ejecución de pruebas de las reglas de la sala de espera')
setp(41,'Se registraron 14 casos aprobados de decisiones locales de la sala de espera. La transmisión de una videollamada no fue comprobada.')
setp(45,'Se registraron 19 casos aprobados de validación clínica, preparación del prediagnóstico y contenido HTML de recetas. No se comprobó la generación final del PDF.')
setp(47,'Figura 3.8. Ejecución de pruebas de pagos simulados')
setp(49,'Se registraron 8 casos aprobados del servicio de simulación de pagos con Supabase sustituido. No se comprobaron pagos reales.')
setp(53,'Se registraron 5 casos aprobados de cálculos locales. El conteo de atenciones completadas fue corregido; el significado de la estimación monetaria permanece pendiente en H-08.')
q=ps[54]
for s,model in [('Figura 3.10. Ejecución de pruebas de componentes aislados',ps[51]),('[Espacio reservado para captura real pendiente de incorporación]',ps[52]),('Se registraron 6 casos adicionales aprobados: tres de pagos simulados, dos de sala de espera y uno de prediagnóstico. Se sustituyeron la navegación y Supabase; no se recorrió la aplicación real.',ps[53]),('Fuente: Elaboración propia, 2026',ps[54])]:q=after(q,s,model)
setp(55,'El apéndice presenta los 114 casos por caso de uso y sus límites. Los pagos reales, la videollamada, la autenticación real y la integración con Supabase no fueron comprobados mediante estas pruebas unitarias. H-03, la parte no resuelta de H-08 y la disponibilidad médica permanecen pendientes.')
setp(58,'Cada fila corresponde a una prueba automatizada de una función o un componente aislado. Jest o unittest prepararon entradas ficticias, ejecutaron el código y comprobaron sus respuestas. Los eventos de componentes se simularon en memoria; no se realizaron recorridos manuales de pantalla. «Éxito» significa que se cumplieron las aserciones, sin implicar ausencia de problemas en el sistema.')
setp(59,'Se conservan los identificadores PU-01 a PU-69 y se incorporan PU-70 a PU-114 para los 45 casos adicionales. Los nuevos identificadores son referencias documentales de las pruebas de regresión. Los estados Pendiente, Requiere integración y No implementado no forman parte de las 114 pruebas ejecutadas ni equivalen a pruebas omitidas.')
after(ps[59],'La ejecución final se realizó el 30 de septiembre de 2026 en Windows, con Node.js 24.15.0 y Python 3.12.14, sobre una versión con cambios locales. Los comandos, la referencia de versión y las fuentes se conservan en ejecucion.json; los resultados están en jest.json, jest.log, python.json y python.log de evidencias/revision-hallazgos-20260930/final. Los registros anteriores se mantienen como respaldo del comportamiento previo.',ps[58])
setp(82,'Los casos de usuarios se relacionan con las figuras 3.2 y 3.3. PU-21 y PU-25 ahora comprueban el rechazo de los campos formados solo por espacios; sus resultados anteriores se conservan en las evidencias históricas.')
setp(102,'Se comprobaron el filtro local de especialidad, la prioridad pediátrica y la ausencia de clasificación inventada cuando falta configuración. H-03 requiere definir qué hacer con triajes sin especialidad. No se midió exactitud clínica ni disponibilidad real de especialistas.')
setp(107,'Se probaron las decisiones locales de cierre y habilitación, además de dos casos de eventos simulados en el componente aislado (figura 3.10). La cola, su actualización real y la comunicación con Supabase requieren integración.')
setp(120,'Se evaluaron la preparación del prediagnóstico y un componente aislado sin triaje (figura 3.10). Se comprobó que no inventa paciente ni urgencia y bloquea la atención sin los datos requeridos. No se comprobó validez clínica ni integración real.')
setp(141,'La verificación bancaria real no está implementada. Los once casos siguientes comprueban ocho respuestas del servicio de simulación y tres del componente aislado de pago, con Supabase y navegación sustituidos (figuras 3.8 y 3.10).')
for p in ps:
    if p.text.startswith('Evidencia:'):
        txt(p,'Evidencia: evidencias/revision-hallazgos-20260930/final. Los resultados de cada prueba se conservan en jest.json o python.json, junto con los registros de ejecución y las fuentes utilizadas.')
findings=[
'H-01. Corregido en la validación local. El registro rechaza el nombre y otros campos personales obligatorios formados solo por espacios. PU-21 y PU-70 a PU-75 comprueban el rechazo y la aceptación de contenido con espacios exteriores.',
'H-02. Corregido en la validación local. El perfil del paciente rechaza teléfono o dirección formados solo por espacios. PU-25 y PU-76 a PU-78 comprueban estos límites.',
'H-03. Pendiente de definición de una regla del sistema. El filtro conserva un triaje sin especialidad al seleccionar Pediatría. PU-37 reproduce ese comportamiento; su aprobación no demuestra una derivación adecuada. Debe decidirse cómo asignar esos triajes.',
'H-04. Corregido en las decisiones y componentes aislados evaluados. Una cita cancelada o sin estado válido no habilita la navegación; el cierre del triaje y una cancelación posterior la bloquean. Lo comprueban PU-51, PU-79 a PU-85, PU-113 y PU-114. No acredita una videollamada real ni corrige el botón de disponibilidad médica.',
'H-05. Corregido en la validación local del historial. Se rechazan diagnóstico y tratamiento en blanco y se conserva la aceptación del contenido con espacios exteriores. Lo comprueban PU-55 y PU-86 a PU-88.',
'H-06. Corregido en la preparación y el componente aislado del prediagnóstico. La ausencia o invalidez de urgencia ya no produce una clasificación inventada ni datos de demostración. Lo comprueban PU-57, PU-58, PU-89 a PU-95 y PU-112.',
'H-07. Corregido en la simulación de pago evaluada. Se rechazan identificadores inválidos y se propagan los errores. El componente aislado evita anunciar éxito o navegar si falla la actualización simulada. Lo comprueban PU-61, PU-63, PU-64, PU-96 a PU-99 y PU-109 a PU-111. No se verificó ninguna transacción bancaria.',
'H-08. Parcialmente corregido y pendiente. PU-67 comprueba que waiting e in_progress no se cuentan como atenciones completadas. La estimación monetaria sigue calculándose a partir de determinados estados multiplicados por 50 y no representa ingresos verificados. Falta definir su significado y ajustar su presentación; PU-67 a PU-69 no acreditan cobros reales.',
'H-09. Corregido en el servicio de análisis evaluado. Se rechazan respuestas incompletas, valores no admitidos y fallos de red o configuración mediante un error que solicita validación manual. Lo comprueban PU-39, PU-41, PU-42 y PU-100 a PU-108. El endpoint FastAPI no fue probado mediante HTTP y no se implementó un flujo nuevo de atención manual.'
]
for i,s in enumerate(findings,159):setp(i,s)
setp(168,'La disponibilidad médica permanece pendiente: el fallo informado del botón no fue reproducido ni se declara corregido. Los pagos reales, la videollamada, la autenticación real y la integración con Supabase no fueron comprobados mediante pruebas unitarias. Tampoco se evaluó precisión clínica ni cobertura completa de líneas o ramas. Las 114 aprobaciones se limitan a las aserciones documentadas.')
cell(d.tables[1].rows[0].cells[1],'Comprobar las reglas locales y las correcciones mediante entradas ficticias, identificando los límites que requieren decisiones del sistema o pruebas de integración.')
cell(d.tables[1].rows[1].cells[1],'Ejecutar 114 pruebas: usuarios (39), triaje (23), decisiones y componentes de sala de espera (16), historial, prediagnóstico y recetas (20), pagos simulados (11) y reportes (5). Se utilizaron sustitutos de los servicios externos.')
cell(d.tables[1].rows[2].cells[1],'Las 114 pruebas aprobaron: 98 de Jest y 16 de unittest, sin fallos ni omisiones. Se comprobaron las correcciones locales descritas en el apéndice. H-03, parte de H-08 y la disponibilidad médica siguen pendientes; la aprobación no acredita los servicios reales ni los módulos completos.')
renderer=read(ROOT/'node_modules/react-test-renderer/package.json')['version']
newrow=d.tables[0].add_row()
for c,s in zip(newrow.cells,['react-test-renderer',renderer,'Monta componentes aislados para comprobar sus respuestas con navegación y servicios sustituidos.']):cell(c,s)

# Preserve each original row, changing only expectations superseded by verified fixes.
rows={}
for ti,t in enumerate(d.tables[2:],2):
    for r in t.rows[1:]:
        n=int(r.cells[0].text.split('-')[1]);rows[n]=(ti,r)
fix={21:'La validación rechaza el registro y solicita completar todos los campos comunes.',25:'La validación rechaza los datos y solicita completar el número de teléfono y la dirección.',39:'El análisis genera un error que solicita validación manual; no inventa una clasificación.',41:'El servicio rechaza la respuesta incompleta con un error que solicita validación manual.',42:'El servicio rechaza la respuesta no interpretable con un error que solicita validación manual.',51:'La decisión local no habilita la llamada para la cita cancelada.',55:'La validación rechaza los campos clínicos formados solo por espacios.',57:'El nombre es «Desconocido» y la urgencia queda sin clasificación (null).',58:'Se obtiene «Desconocido», urgencia sin clasificación (null) y «Sin síntomas reportados», sin datos de demostración.',61:'Se genera el error «Identificador de triaje requerido» y no se solicita una actualización al sustituto de Supabase.',63:'La promesa rechaza con el error ficticio devuelto; el error no se oculta ni se limita a registrarlo.',64:'La promesa propaga la misma excepción ficticia de la actualización.',67:'El contador de atenciones completadas es cero. La estimación monetaria sigue siendo 100; no representa ingresos comprobados.'}
for n,s in fix.items():
    r=rows[n][1];cell(r.cells[3],s);cell(r.cells[4],'Éxito')
for n in [21,25,39,41,42,51,55,57,58,61,63,64,67]:
    r=rows[n][1]
    if n in [39,41,42,63,64]:cell(r.cells[2],r.cells[2].text.split('\n3.')[0]+'\n3. Comprobar el rechazo y el error esperado.')
cell(rows[49][1].cells[2],'1. Preparar un triaje waiting y una cita ficticia con estado scheduled.\n2. Invocar la decisión local de habilitación.\n3. Comprobar que devuelve verdadero.')
cell(rows[63][1].cells[2],'1. Preparar el identificador «triaje-ficticio» y una respuesta simulada con «Fallo ficticio».\n2. Ejecutar la simulación.\n3. Comprobar el rechazo con ese error y que no se limita a registrarlo en consola.')
cell(rows[64][1].cells[2],'1. Preparar el identificador «triaje-ficticio» y una excepción «Fallo ficticio».\n2. Hacer que la actualización simulada lance la excepción.\n3. Comprobar que la función propaga la misma excepción.')
cell(rows[37][1].cells[4],'Éxito como caracterización; regla de asignación pendiente.')
for n in [67,68,69]:cell(rows[n][1].cells[4],'Éxito en Jest; significado monetario pendiente (H-08).')

assertions=[]
for suite in j['testResults']:
    for a in suite['assertionResults']:assertions.append({'file':Path(suite['name']).name,'title':a['title'],'fullName':a['fullName'],'status':a['status'],'framework':'Jest'})
for a in py['cases']:assertions.append({'file':'test_unit_llm.py','title':a['id'],'fullName':a['id'],'status':a['status'],'framework':'unittest'})
mapping=[];used=set()
for n in range(1,70):
    if n<=5:a=[a for a in assertions if a['file']=='loginValidation.test.ts'][n-1]
    else:a=next(a for a in assertions if re.search(r'PU[-_]0?'+str(n)+r'(?!\d)',a['title']))
    used.add(a['fullName']);mapping.append({'id':f'PU-{n:02}',**a})
remaining=lambda f:[a for a in assertions if a['file']==f and a['fullName'] not in used]
def add(a,ti,scenario,steps,expected):
    n=len(mapping)+1;t=d.tables[ti];el=copy.deepcopy(t.rows[1]._tr);t._tbl.append(el);r=t.rows[-1]
    for c,s in zip(r.cells,[f'PU-{n:02}',scenario,steps,expected,'Éxito']):cell(c,s)
    for c in r.cells:
        for p in c.paragraphs:p.paragraph_format.keep_with_next=False
    mapping.append({'id':f'PU-{n:02}',**a});used.add(a['fullName'])
def steps(data,action):return f'1. {data}\n2. {action}\n3. Comparar la respuesta con el resultado esperado.'
a=remaining('userValidation.test.ts')
for x,field in zip(a[:4],['correo','carnet de identidad','teléfono','dirección']):add(x,3,f'El campo de {field} contiene solo espacios.',steps(f'Preparar un registro ficticio completo y sustituir {field} por espacios, tabulación y salto de línea.','Invocar la validación del registro.'),'La validación rechaza el registro.')
add(a[4],3,'El nombre contiene espacios exteriores.',steps('Preparar el registro ficticio con « Persona de Prueba ».','Invocar la validación del registro.'),'El registro supera la validación local.')
add(a[5],3,'La especialidad del médico contiene solo espacios.',steps('Preparar un médico ficticio completo con especialidad de tres espacios.','Invocar la validación del registro.'),'La validación rechaza el registro.')
for x,data in zip(a[6:8],['Usar teléfono de espacios y tabulación, con «Dirección ficticia».','Usar teléfono «00000000» y dirección de salto de línea y espacios.']):add(x,4,'Un campo del perfil está en blanco y el otro tiene contenido.',steps(data,'Invocar la validación del perfil de paciente.'),'La validación rechaza los datos.')
add(a[8],4,'Los campos completos tienen espacios exteriores.',steps('Usar « 00000000 » y « Dirección ficticia ».','Invocar la validación del perfil de paciente.'),'Los datos superan la validación local.')
a=remaining('consultationRules.test.ts')
for x,status in zip(a[:4],['cancelled','completed','unknown','sin definir']):add(x,9,f'La cita tiene estado {status} y el triaje sigue en progreso.',steps(f'Preparar triaje in_progress y cita con estado {status}.','Invocar la decisión de habilitación.'),'La función no habilita la llamada.')
for x,status in zip(a[4:6],['completed','resolved']):add(x,9,f'El triaje está cerrado como {status}.',steps(f'Preparar un triaje {status} y una cita scheduled.','Invocar la decisión de habilitación.'),'La función no habilita la llamada.')
add(a[6],9,'La cita tiene identificador, pero no tiene estado.',steps('Preparar triaje waiting y cita con identificador ficticio sin estado.','Invocar la decisión de habilitación.'),'La función no habilita la llamada.')
a=remaining('medicalRecord.test.ts')
for x,data in zip(a[:2],['Usar diagnóstico de espacios y tabulación, y «Plan ficticio».','Usar «Diagnóstico ficticio» y tratamiento de salto de línea y espacios.']):add(x,11,'Un campo clínico está en blanco y el otro tiene contenido.',steps(data,'Invocar la validación clínica.'),'La validación rechaza los campos.')
add(a[2],11,'El contenido clínico tiene espacios exteriores.',steps('Usar « Diagnóstico ficticio » y « Plan ficticio ».','Invocar la validación clínica.'),'La validación acepta el contenido.')
for x,level in zip(a[3:6],['Critical','Medium','Low']):add(x,10,f'El reporte contiene urgencia {level}.',steps(f'Preparar un reporte ficticio con urgencia {level}.','Invocar la preparación del prediagnóstico.'),f'Se conserva la urgencia {level} recibida.')
for x,level in zip(a[6:],['cadena vacía','Unknown','tres espacios','null']):add(x,10,f'El reporte tiene un nivel no admitido: {level}.',steps(f'Preparar un reporte con urgencia igual a {level}.','Invocar la preparación del prediagnóstico.'),'La urgencia queda sin clasificación (null).')
for x,val in zip(remaining('paymentSimulation.test.ts'),['cadena vacía','tres espacios','lista con «uno» y «dos»','lista vacía']):add(x,13,f'El identificador de la simulación es una {val}.' if val.startswith('lista') else f'El identificador contiene {val}.',steps(f'Preparar {val} como identificador, con Supabase sustituido.','Ejecutar el servicio de simulación.'),'Se rechaza con «Identificador de triaje requerido» y no se solicita actualización.')
for x in remaining('test_unit_llm.py'):
    key=x['title'].split('test_H09_')[1]
    values={'fallo_red_no_inventa_urgencia':'Configurar la red simulada para lanzar «Fallo ficticio de red».','rechaza_especialidad_ausente':'Usar una respuesta completa y sustituir la especialidad por un valor ausente.','rechaza_especialidad_desconocida':'Usar una respuesta completa y especialidad «Especialidad inexistente».','rechaza_idioma_ausente':'Usar una respuesta completa y sustituir el idioma por un valor ausente.','rechaza_recomendacion_ausente':'Usar una respuesta completa y sustituir la recomendación por un valor ausente.','rechaza_sintomas_blancos':'Usar una respuesta completa y síntomas de espacios y tabulación.','rechaza_sintomas_tipo_incorrecto':'Usar una respuesta completa y el número 5 como síntomas.','rechaza_urgencia_ausente':'Usar una respuesta completa y sustituir la urgencia por un valor ausente.','rechaza_urgencia_desconocida':'Usar una respuesta completa y urgencia «Unknown».'}
    scenarios={'fallo_red_no_inventa_urgencia':'El proveedor simulado produce un fallo de red.','rechaza_especialidad_ausente':'La respuesta no contiene especialidad.','rechaza_especialidad_desconocida':'La respuesta contiene una especialidad no admitida.','rechaza_idioma_ausente':'La respuesta no contiene idioma.','rechaza_recomendacion_ausente':'La respuesta no contiene recomendación.','rechaza_sintomas_blancos':'Los síntomas de la respuesta contienen solo espacios.','rechaza_sintomas_tipo_incorrecto':'Los síntomas de la respuesta son un número.','rechaza_urgencia_ausente':'La respuesta no contiene urgencia.','rechaza_urgencia_desconocida':'La respuesta contiene una urgencia no admitida.'}
    add(x,7,scenarios[key],steps(values[key],'Invocar el análisis con «relato ficticio» y el proveedor sustituido.'),'El servicio genera un error que solicita validación manual y no inventa una clasificación.')
a=remaining('reviewScreens.test.tsx')
for x,mode in zip(a[:2],['respuesta con error','excepción']):add(x,13,f'La actualización simulada produce una {mode}.',steps(f'Montar el componente de pago con identificador ficticio y {mode} en el sustituto.','Invocar su evento de pago y avanzar los temporizadores simulados.'),'Muestra el fallo, permite reintentar y no anuncia éxito ni solicita navegación.')
add(a[2],13,'La actualización simulada termina sin error.',steps('Montar el componente de pago y simular una actualización sin error.','Invocar el evento de pago y avanzar los temporizadores.'),'Muestra «Simulación completada» y solicita navegar a la sala de espera; no afirma validar una transferencia.')
add(a[3],10,'El componente no recibe un triaje.',steps('Montar el componente de prediagnóstico sin triaje.','Inspeccionar el texto y los controles del componente aislado.'),'Muestra «URGENCIA NO DISPONIBLE», no muestra paciente de demostración ni nivel crítico y bloquea atención y navegación.')
add(a[4],9,'Llegan eventos simulados para una cita cancelada.',steps('Montar la sala de espera con cita cancelled y servicios sustituidos.','Emitir los eventos simulados de cita, triaje y aviso del médico; avanzar temporizadores.'),'No se solicita navegar a la videoconsulta.')
add(a[5],9,'Se cancela una cita antes de la navegación prevista.',steps('Montar la sala de espera con cita scheduled.','Cambiarla a cancelled, emitir el evento simulado y avanzar temporizadores.'),'La navegación prevista queda revocada y no se solicita acceder a la videoconsulta.')
assert len(mapping)==114 and len(used)==114
assert all(x['status']=='passed' for x in mapping)
assert sum(x['framework']=='Jest' for x in mapping)==98
assert len(d.tables)==15
ps[17]._p.addprevious(ps[55]._p)
for r in d.tables[1].rows:
    for c in r.cells:
        for p in c.paragraphs:p.paragraph_format.keep_with_next=True
for i,p in enumerate(d.paragraphs[:-1]):
    if p.text.startswith('Caso de uso '):
        following=d.paragraphs[i+1]
        following.paragraph_format.keep_with_next=bool(i+2<len(d.paragraphs) and d.paragraphs[i+2].text.startswith('Tabla A.'))
for t in d.tables[2:]:
    for c in t.rows[-1].cells:
        for p in c.paragraphs:p.paragraph_format.keep_with_next=True
    nxt=t._tbl.getnext()
    if nxt is not None and nxt.tag.endswith('}p'):Paragraph(nxt,t._parent).paragraph_format.keep_with_next=True
# Give every placeholder its figure identifier, retaining the original visual style.
fig=None
for p in d.paragraphs:
    if p.text.startswith('Figura '):fig=p.text.split('. ',1)[0]
    if p.text.startswith('[Espacio reservado'):txt(p,f'[Captura pendiente de incorporación — {fig}]')
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                for run in p.runs:run.font.name='Arial';run.font.size=Pt(10.5)
d.save(OUT)
assert sha(SRC)==original
(QA/'casos-documentados.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'auditoria.json').write_text(json.dumps({'original_sha256':original,'original_preservado':True,'pruebas':114,'Jest':98,'unittest':16,'fallidas':0,'omitidas':0,'tablas_capitulo':2,'tablas_apendice':13,'capturas_pendientes':16,'fuentes_coinciden_con_ejecucion':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT)
